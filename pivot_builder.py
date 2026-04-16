"""
Сводная таблица: Компонент × Показатель
========================================
Автоматически определяет все показатели из CSV-файла.

Зависимости:
    pip install pandas openpyxl

Использование:
    python pivot_builder.py
    python pivot_builder.py --input data.csv --output result.xlsx
"""

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─── Настройки колонок ────────────────────────────────────────────────────────

# Имена колонок в исходном файле (скрипт попробует найти их автоматически,
# если точные названия не совпадут — отредактируйте здесь)
COL_COMPONENT = "Компонент"
COL_BATCH     = "Наименование партии"
COL_INDICATOR = "Наименование показателя"
COL_UNIT      = "Единица измерения_по_партиям"
COL_VALUE     = "Значение показателя"

# ─── Цвета ────────────────────────────────────────────────────────────────────

C_HEADER      = "1F4E79"   # тёмно-синий — шапка
C_INDEX       = "2E75B6"   # средне-синий — колонки компонент/партия чётные
C_INDEX_ALT   = "D6E4F0"   # светло-синий — нечётные
C_DATA_ALT    = "EBF3FB"   # очень светло-синий — данные чётные
C_DATA        = "FFFFFF"   # белый — данные нечётные
C_MISSING     = "FFE0E0"   # розовый — пропуск
C_WARN_FILL   = "FFF2CC"   # жёлтый — предупреждение в шапке анализа


def auto_detect_columns(df: pd.DataFrame) -> dict:
    """Пытается угадать нужные колонки по ключевым словам."""
    mapping = {}
    targets = {
        COL_COMPONENT: ["компонент", "component"],
        COL_BATCH:     ["парти", "batch", "lot"],
        COL_INDICATOR: ["показател", "indicator", "property", "параметр"],
        COL_UNIT:      ["единиц", "unit", "мер"],
        COL_VALUE:     ["значени", "value", "val"],
    }
    cols_lower = {c.lower(): c for c in df.columns}
    for target, keywords in targets.items():
        for col_l, col_orig in cols_lower.items():
            if any(kw in col_l for kw in keywords):
                mapping[target] = col_orig
                break
    return mapping


def build_pivot(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Строит pivot и две вспомогательные таблицы с анализом пропусков."""

    # Очистка значений (запятая как десятичный разделитель)
    df[COL_VALUE] = (
        df[COL_VALUE]
        .astype(str)
        .str.replace(",", ".", regex=False)
        .str.strip()
    )
    df[COL_VALUE] = pd.to_numeric(df[COL_VALUE], errors="coerce")

    df_clean = df.dropna(subset=[COL_INDICATOR]).copy()

    pivot = df_clean.pivot_table(
        index=[COL_COMPONENT, COL_BATCH],
        columns=COL_INDICATOR,
        values=COL_VALUE,
        aggfunc="first",
    )
    pivot.columns.name = None  # убираем label "Наименование показателя"

    n_rows, n_cols = pivot.shape

    # --- Анализ пропусков по строкам ---
    miss_row = pivot.isnull().sum(axis=1)
    df_miss_rows = pd.DataFrame({
        COL_COMPONENT:     [i[0] for i in miss_row.index],
        COL_BATCH:         [i[1] for i in miss_row.index],
        "Пропусков":       miss_row.values,
        "Всего показателей": n_cols,
        "Заполнено, %":    ((n_cols - miss_row.values) / n_cols * 100).round(1),
    }).sort_values("Пропусков", ascending=False).reset_index(drop=True)

    # --- Анализ пропусков по столбцам ---
    miss_col = pivot.isnull().sum(axis=0)
    df_miss_cols = pd.DataFrame({
        "Показатель":      miss_col.index,
        "Пропусков":       miss_col.values,
        "Всего партий":    n_rows,
        "Заполнено, %":    ((n_rows - miss_col.values) / n_rows * 100).round(1),
    }).sort_values("Пропусков", ascending=False).reset_index(drop=True)

    return pivot, df_miss_rows, df_miss_cols


def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def fmt_header(cell, *, wrap=True):
    cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=9)
    cell.fill      = PatternFill("solid", start_color=C_HEADER)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    cell.border    = thin_border()


def fmt_index(cell, even: bool):
    cell.font      = Font(bold=True, name="Arial", size=9)
    cell.fill      = PatternFill("solid", start_color=C_INDEX if even else C_INDEX_ALT)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border    = thin_border()


def fmt_data(cell, even: bool, missing: bool):
    cell.font      = Font(name="Arial", size=9)
    cell.fill      = PatternFill("solid", start_color=(
        C_MISSING if missing else (C_DATA_ALT if even else C_DATA)
    ))
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = thin_border()


def fmt_analysis(cell, even: bool):
    cell.font      = Font(name="Arial", size=9)
    cell.fill      = PatternFill("solid", start_color=C_DATA_ALT if even else C_DATA)
    cell.alignment = Alignment(
        horizontal="left" if cell.column == 1 else "center",
        vertical="center",
    )
    cell.border    = thin_border()


def apply_formatting(output_path: str):
    wb = load_workbook(output_path)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Шапка
        for cell in ws[1]:
            fmt_header(cell)
        ws.row_dimensions[1].height = 55

        if sheet_name == "Сводная таблица":
            # Индексные колонки (Компонент + Партия = первые 2)
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
                for cell in row:
                    fmt_index(cell, even=(cell.row % 2 == 0))

            # Данные
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=ws.max_column):
                for cell in row:
                    fmt_data(cell, even=(cell.row % 2 == 0), missing=(cell.value is None))

            # Ширины
            ws.column_dimensions["A"].width = 30
            ws.column_dimensions["B"].width = 18
            for col_idx in range(3, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = 15

            ws.freeze_panes = "C2"

        else:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    fmt_analysis(cell, even=(cell.row % 2 == 0))

            # Авто-ширина
            for col in ws.columns:
                max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 55)

    wb.save(output_path)


def print_summary(pivot: pd.DataFrame, df_miss_rows: pd.DataFrame, df_miss_cols: pd.DataFrame):
    n_rows, n_cols = pivot.shape
    total   = n_rows * n_cols
    missing = pivot.isnull().sum().sum()

    print(f"\n{'─'*50}")
    print(f"  Компонент × Показатель: {n_rows} × {n_cols}")
    print(f"  Пропуски: {missing}/{total}  ({100*missing/total:.1f}%)")
    print(f"{'─'*50}")

    print("\nТоп-5 компонентов с наибольшим числом пропусков:")
    for _, r in df_miss_rows.head(5).iterrows():
        bar = "█" * int((r["Пропусков"] / n_cols) * 20)
        print(f"  {r[COL_COMPONENT]:<30} {r[COL_BATCH]:<15}  {bar}  {r['Пропусков']}/{n_cols}")

    print("\nТоп-5 показателей с наибольшим числом пропусков:")
    for _, r in df_miss_cols.head(5).iterrows():
        bar = "█" * int((r["Пропусков"] / n_rows) * 20)
        print(f"  {r['Показатель']:<45}  {bar}  {r['Пропусков']}/{n_rows}")
    print()


def save_csv(pivot: pd.DataFrame, df_miss_rows: pd.DataFrame,
             df_miss_cols: pd.DataFrame, base_path: Path, encoding: str):
    """Сохраняет три CSV-файла рядом с основным выходным файлом."""
    stem = base_path.stem
    out_dir = base_path.parent

    paths = {
        "сводная таблица":          out_dir / f"{stem}_pivot.csv",
        "пропуски по компонентам":  out_dir / f"{stem}_missing_components.csv",
        "пропуски по показателям":  out_dir / f"{stem}_missing_indicators.csv",
    }

    pivot.to_csv(paths["сводная таблица"], encoding=encoding)
    df_miss_rows.to_csv(paths["пропуски по компонентам"], index=False, encoding=encoding)
    df_miss_cols.to_csv(paths["пропуски по показателям"], index=False, encoding=encoding)

    for label, path in paths.items():
        print(f"  CSV [{label}]: {path.resolve()}")


def main():
    parser = argparse.ArgumentParser(description="Сводная таблица компонент × показатель")
    parser.add_argument("--input",    default="daimler_component_properties.csv",
                        help="Путь к исходному CSV-файлу")
    parser.add_argument("--output",   default="pivot_components.xlsx",
                        help="Путь к выходному Excel-файлу")
    parser.add_argument("--encoding", default="utf-8-sig",
                        help="Кодировка CSV (по умолчанию utf-8-sig)")
    parser.add_argument("--no-excel", action="store_true",
                        help="Не создавать Excel-файл, только CSV")
    args = parser.parse_args()

    input_path  = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"[ОШИБКА] Файл не найден: {input_path}")
        sys.exit(1)

    print(f"Читаю: {input_path}")
    df = pd.read_csv(input_path, encoding=args.encoding)

    # Автоопределение колонок
    detected = auto_detect_columns(df)
    missing_cols = [c for c in [COL_COMPONENT, COL_BATCH, COL_INDICATOR, COL_VALUE]
                    if c not in df.columns]
    if missing_cols:
        print(f"\n[!] Колонки не найдены точно: {missing_cols}")
        print(f"    Автоопределение: {detected}")
        print(f"    Доступные колонки: {list(df.columns)}")
        for expected, found in detected.items():
            if expected in missing_cols and found:
                df.rename(columns={found: expected}, inplace=True)
                print(f"    Переименовано: '{found}' → '{expected}'")

    print(f"Показателей обнаружено: {df[COL_INDICATOR].nunique()}")
    print(f"Компонентов обнаружено: {df[COL_COMPONENT].nunique()}")
    print(f"Партий обнаружено:      {df[COL_BATCH].nunique()}")

    pivot, df_miss_rows, df_miss_cols = build_pivot(df)

    # Сохраняем CSV
    print("\nСохраняю CSV:")
    save_csv(pivot, df_miss_rows, df_miss_cols, output_path, args.encoding)

    # Сохраняем Excel (если не отключён)
    if not args.no_excel:
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            pivot.to_excel(writer, sheet_name="Сводная таблица")
            df_miss_rows.to_excel(writer, sheet_name="Пропуски по компонентам", index=False)
            df_miss_cols.to_excel(writer, sheet_name="Пропуски по показателям", index=False)

        apply_formatting(str(output_path))
        print(f"\nСохраняю Excel:")
        print(f"  XLSX: {output_path.resolve()}")

    print_summary(pivot, df_miss_rows, df_miss_cols)


if __name__ == "__main__":
    main()
